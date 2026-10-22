"""Regenerate the paper's six figures from attachments and solver evidence.

The generator reads the official attachments and the paper-facing result/table
files.  It does not contain paper-result constants (apart from labels and
physical limits), so a fresh run is a data-to-figure check rather than a copy
of the existing PNGs.

Run after generating the annual evidence::

    python code/run.py figure.generate_paper_figures

Set MCM_OUTPUT_ROOT to select the annual evidence; --output-dir selects only
the generated image destination inside project output/.
"""

from __future__ import annotations

import argparse
from datetime import datetime, time
from pathlib import Path
from typing import Any
from ..common.paths import output_root, protected_output_path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
PROJECT_DIR = ROOT / "data"
ATTACHMENT_DIR = PROJECT_DIR / "附件"
RESULTS_DIR = output_root() / "result" / "来源证据"
DEFAULT_OUTPUT_DIR = output_root() / "figure"

INTERVALS = 144
INTERVAL_HOURS = 1.0 / 6.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
TARGET_DATES = (
    ("2025-03-20", "春分 3/20"),
    ("2025-06-21", "夏至 6/21"),
    ("2025-09-23", "秋分 9/23"),
    ("2025-12-21", "冬至 12/21"),
)

COLORS = {
    "load": "#1f4e79",
    "pv": "#d97706",
    "pv_fill": "#fbbf24",
    "net": "#5b4b8a",
    "price": "#c2410c",
    "charge": "#0f766e",
    "discharge": "#65a30d",
    "emergency": "#c62828",
    "soc": "#4c1d95",
    "bound": "#b91c1c",
    "cost": "#1f4e79",
    "text": "#1f2937",
    "grid": "#e5e7eb",
    "mute": "#6b7280",
}


def setup_style() -> None:
    """Use a stable, paper-compatible plotting style."""

    plt.rcParams.update(
        {
            "font.sans-serif": [
                "Microsoft YaHei",
                "SimHei",
                "Arial Unicode MS",
                "Noto Sans CJK SC",
                "DejaVu Sans",
            ],
            "axes.unicode_minus": False,
            "axes.edgecolor": "#9ca3af",
            "axes.labelcolor": COLORS["text"],
            "xtick.color": COLORS["text"],
            "ytick.color": COLORS["text"],
            "axes.linewidth": 0.8,
            "axes.labelsize": 10.5,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "legend.fontsize": 9,
            "grid.color": COLORS["grid"],
            "grid.linewidth": 0.7,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def polish_axes(ax: Any, *, grid_y: bool = True) -> None:
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    if grid_y:
        ax.grid(axis="y", alpha=0.9, zorder=0)
    ax.set_axisbelow(True)


def save_figure(fig: Any, output_dir: Path, filename: str, *, dpi: int) -> Path:
    output_dir = protected_output_path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / filename
    fig.savefig(path, dpi=dpi, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print(path)
    return path


def _as_float(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"expected numeric value, got {value!r}") from error


def _time_to_minutes(value: Any) -> int:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip().replace("：", ":")
    if text.endswith("+1"):
        text = text[:-2]
    hour, minute = (int(part) for part in text.split(":")[:2])
    return hour * 60 + minute


def _interval_grid() -> np.ndarray:
    return np.arange(INTERVALS, dtype=float) * INTERVAL_HOURS


def _step_grid(values: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    if len(values) != INTERVALS:
        raise ValueError(f"expected {INTERVALS} intervals, got {len(values)}")
    return np.arange(INTERVALS + 1, dtype=float) * INTERVAL_HOURS, np.r_[values, values[-1]]


def _read_attachment1() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    path = ATTACHMENT_DIR / "附件1.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    if len(rows) != INTERVALS + 1 or len(rows[0]) < 4:
        raise ValueError(f"附件1 shape is not 144 x 4: {path}")
    data = rows[1:]
    x = _interval_grid()
    price = np.asarray([_as_float(row[1]) for row in data], dtype=float)
    load = np.asarray([_as_float(row[2]) for row in data], dtype=float)
    pv = np.asarray([_as_float(row[3]) for row in data], dtype=float)
    return x, price, load, pv


def _read_wide_attachment(path: Path, sheet_name: str) -> tuple[pd.DatetimeIndex, np.ndarray]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"missing sheet {sheet_name!r} in {path}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if len(rows) != 366 or len(rows[0]) < INTERVALS + 1:
        raise ValueError(f"unexpected shape for {path}/{sheet_name}: {len(rows)} rows")
    dates = pd.DatetimeIndex(pd.to_datetime([row[0] for row in rows[1:]], errors="raise"))
    matrix = np.asarray(
        [[_as_float(value) for value in row[1 : INTERVALS + 1]] for row in rows[1:]],
        dtype=float,
    )
    if matrix.shape != (365, INTERVALS):
        raise ValueError(f"unexpected data shape for {path}/{sheet_name}: {matrix.shape}")
    return dates, matrix


def _read_attachment2_day(date_text: str) -> tuple[np.ndarray, np.ndarray]:
    path = ATTACHMENT_DIR / "附件2.xlsx"
    load_dates, load_matrix = _read_wide_attachment(path, "小区负载")
    pv_dates, pv_matrix = _read_wide_attachment(path, "光伏发电实际功率")
    target = pd.Timestamp(date_text)
    load_index = np.flatnonzero(load_dates == target)
    pv_index = np.flatnonzero(pv_dates == target)
    if len(load_index) != 1 or len(pv_index) != 1:
        raise ValueError(f"附件2 missing unique date {date_text}")
    return load_matrix[load_index[0]], pv_matrix[pv_index[0]]


def _read_attachment3_day(date_text: str) -> dict[float, np.ndarray]:
    path = ATTACHMENT_DIR / "附件3.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    target = pd.Timestamp(date_text).date()
    current_date = None
    releases: dict[float, np.ndarray] = {}
    for row in rows[1:]:
        if row[0] not in (None, ""):
            current_date = pd.Timestamp(str(row[0]).strip()).date()
        if current_date != target or row[1] in (None, ""):
            continue
        release_hour = _time_to_minutes(row[1]) / 60.0
        values = np.asarray([_as_float(value) for value in row[2 : 2 + 24]], dtype=float)
        if len(values) != 24:
            raise ValueError(f"附件3 forecast row has {len(values)} values on {date_text}")
        releases[release_hour] = values
    if set(releases) != {0.0, 6.0, 12.0, 18.0}:
        raise ValueError(f"附件3 releases for {date_text}: {sorted(releases)}")
    return releases


def _read_p1_schedule() -> pd.DataFrame:
    path = RESULTS_DIR / "problem1" / "audit" / "problem1_schedule-A.csv"
    frame = pd.read_csv(path)
    required = {"purchase_kwh", "charge_kwh", "discharge_kwh", "soc_start_kwh", "soc_end_kwh"}
    if len(frame) != INTERVALS or not required.issubset(frame.columns):
        raise ValueError(f"unexpected Problem 1 schedule: {path}")
    return frame


def _find_p2_compare() -> Path:
    path = output_root() / "table" / "p2_compare.csv"
    if not path.is_file():
        raise FileNotFoundError(f"Generate p2_compare.csv from the selected results first: {path}")
    return path


def _draw_box(ax: Any, x: float, y: float, width: float, height: float, text: str, color: str) -> None:
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.02,rounding_size=0.03",
        linewidth=1.0,
        edgecolor=color,
        facecolor="#ffffff",
    )
    ax.add_patch(patch)
    ax.text(x + width / 2, y + height / 2, text, ha="center", va="center", color=COLORS["text"], fontsize=9.5, wrap=True)


def _arrow(ax: Any, start: tuple[float, float], end: tuple[float, float], color: str) -> None:
    ax.add_patch(
        FancyArrowPatch(
            start,
            end,
            arrowstyle="-|>",
            mutation_scale=10,
            linewidth=1.2,
            color=color,
            connectionstyle="arc3,rad=0.0",
        )
    )


def figure1_flowchart(output_dir: Path, dpi: int) -> Path:
    fig, ax = plt.subplots(figsize=(12.5, 7.4), constrained_layout=True)
    ax.set_xlim(0, 4)
    ax.set_ylim(0, 10)
    ax.axis("off")
    columns = [
        ("问题一", "#d9778f", [
            ("建立电量平衡关系", 8.1),
            ("建立储能运行约束", 6.5),
            ("最小化全天计划购电费用", 4.8),
            ("求解混合整数线性规划\n（MILP）", 2.7),
        ]),
        ("问题二", "#c9972b", [
            ("完全信息求解\n（问题一的 MILP）", 8.1),
            ("理论最低购电费用\n作为评价基准", 6.3),
            ("扩展 MILP：预测误差\n与紧急购电", 4.5),
            ("日前购电与储能\n调度方案", 2.7),
        ]),
        ("问题三", "#4ca69a", [
            ("0:00 制定计划购电", 8.1),
            ("6:00 / 12:00 / 18:00\n获得新的光伏预报", 6.2),
            ("锁定已执行决策及原储能方案\n重新优化剩余购电量", 4.1),
            ("仅执行计划 / 采用调整策略\n比较费用与紧急购电量", 1.9),
        ]),
        ("问题四", "#8e73bd", [
            ("P4-2：重做问题二", 8.1),
            ("电价预测\nWD2 / Yest / Att1", 6.3),
            ("P4-3：重做问题三", 4.5),
            ("价格预测 + 光伏预报\n滚动调整与日终结算", 2.7),
        ]),
    ]
    width = 0.82
    for column_index, (title, color, boxes) in enumerate(columns):
        x = column_index + 0.09
        ax.add_patch(
            FancyBboxPatch(
                (x - 0.04, 0.35),
                0.9,
                9.3,
                boxstyle="round,pad=0.01",
                linewidth=1.0,
                edgecolor="#9ca3af",
                facecolor="#ffffff",
                linestyle="--",
            )
        )
        ax.text(x + width / 2, 9.35, title, ha="center", va="center", fontsize=13, color=color, fontweight="bold")
        for text, y in boxes:
            _draw_box(ax, x, y, width, 0.95, text, color)
        for (_, upper_y), (_, lower_y) in zip(boxes, boxes[1:]):
            _arrow(ax, (x + width / 2, upper_y), (x + width / 2, lower_y + 0.95), color)
    _arrow(ax, (0.95, 5.0), (1.05, 5.0), "#c9972b")
    _arrow(ax, (1.95, 5.0), (2.05, 5.0), "#4ca69a")
    _arrow(ax, (2.95, 5.0), (3.05, 5.0), "#8e73bd")
    ax.text(1.0, 5.35, "负载与\n光伏预测", ha="center", va="bottom", fontsize=8, color=COLORS["text"])
    ax.text(2.0, 5.35, "购电调整\n与新预报", ha="center", va="bottom", fontsize=8, color=COLORS["text"])
    ax.text(3.0, 5.35, "波动电价\n与价格预测", ha="center", va="bottom", fontsize=8, color=COLORS["text"])
    return save_figure(fig, output_dir, "图01_微网购电与储能调度的整体建模框架.png", dpi=dpi)


def figure2_typical_day(output_dir: Path, dpi: int) -> Path:
    x, price, load, pv = _read_attachment1()
    net = load - pv
    fig, (top, bottom) = plt.subplots(2, 1, figsize=(11.2, 6.4), sharex=True, constrained_layout=True)
    top.plot(x, load, color=COLORS["load"], lw=2.2, label="小区负荷")
    top.plot(x, pv, color=COLORS["pv"], lw=2.2, label="光伏预测")
    top.fill_between(x, 0, pv, color=COLORS["pv_fill"], alpha=0.28)
    top.set_ylabel("功率/kW")
    top.legend(frameon=False, ncol=2, loc="upper left")
    polish_axes(top)

    bottom.plot(x, net, color=COLORS["net"], lw=2.2, label="净负荷 N=L−PV")
    price_axis = bottom.twinx()
    xs, ps = _step_grid(price)
    price_axis.step(xs, ps, where="post", color=COLORS["price"], ls="--", lw=1.8, label="外网电价")
    bottom.set_ylabel("净负荷/kW", color=COLORS["net"])
    price_axis.set_ylabel("电价/(元/kWh)", color=COLORS["price"])
    bottom.tick_params(axis="y", labelcolor=COLORS["net"])
    price_axis.tick_params(axis="y", labelcolor=COLORS["price"])
    lines = bottom.get_lines() + price_axis.get_lines()
    bottom.legend(lines, [line.get_label() for line in lines], frameon=False, ncol=2, loc="upper left")
    bottom.set_xlabel("日内时间/h")
    polish_axes(bottom)
    price_axis.spines["top"].set_visible(False)
    for axis in (top, bottom):
        axis.set_xlim(0, 24)
        axis.set_xticks(range(0, 25, 4))
    return save_figure(fig, output_dir, "图02_典型日负荷光伏电价净负荷.png", dpi=dpi)


def figure3_problem1_dispatch(output_dir: Path, dpi: int) -> Path:
    _, price, _, _ = _read_attachment1()
    schedule = _read_p1_schedule()
    purchase = schedule["purchase_kwh"].to_numpy(float)
    charge = schedule["charge_kwh"].to_numpy(float)
    discharge = schedule["discharge_kwh"].to_numpy(float)
    soc = np.r_[schedule["soc_start_kwh"].iloc[0], schedule["soc_end_kwh"].to_numpy(float)]
    fig, axes = plt.subplots(3, 1, figsize=(11.2, 8.2), sharex=True, constrained_layout=True)

    axis = axes[0]
    xp, yp = _step_grid(purchase)
    axis.step(xp, yp, where="post", color=COLORS["load"], lw=1.8, label="计划购电量")
    axis.set_ylabel("购电量（kWh）", color=COLORS["load"])
    axis.tick_params(axis="y", labelcolor=COLORS["load"])
    polish_axes(axis)
    price_axis = axis.twinx()
    xc, yc = _step_grid(price)
    price_axis.step(xc, yc, where="post", color=COLORS["price"], lw=1.5, ls="--", label="电价")
    price_axis.set_ylabel("电价（元/kWh）", color=COLORS["price"])
    price_axis.tick_params(axis="y", labelcolor=COLORS["price"])
    price_axis.spines["top"].set_visible(False)
    lines = axis.get_lines() + price_axis.get_lines()
    axis.legend(lines, [line.get_label() for line in lines], loc="upper left", frameon=False)

    axis = axes[1]
    xs = np.arange(INTERVALS + 1, dtype=float) * INTERVAL_HOURS
    axis.fill_between(xs, SOC_MIN, SOC_MAX, color=COLORS["soc"], alpha=0.08)
    axis.plot(xs, soc, color=COLORS["soc"], lw=2.0, label="储电量 $E_t$")
    axis.axhline(SOC_INITIAL, color="#64748b", ls="--", lw=1.1, label="日初/日末")
    axis.axhline(SOC_MIN, color=COLORS["bound"], ls=":", lw=1.2, label="运行下限")
    axis.axhline(SOC_MAX, color=COLORS["charge"], ls=":", lw=1.2, label="运行上限")
    axis.set_ylabel("储电量（kWh）")
    axis.set_ylim(0, 12500)
    axis.legend(loc="upper left", bbox_to_anchor=(1.01, 1.0), frameon=False, fontsize=8)
    polish_axes(axis)

    axis = axes[2]
    xp, yp = _step_grid(charge)
    axis.step(xp, yp, where="post", color=COLORS["charge"], lw=1.7, label="充电量 $u_t$")
    xp, yp = _step_grid(discharge)
    axis.step(xp, yp, where="post", color=COLORS["discharge"], lw=1.7, label="放电量 $v_t$")
    axis.set_ylabel("电量（kWh）")
    axis.set_xlabel("日内时间/h")
    axis.set_xlim(0, 24)
    axis.set_xticks(range(0, 25, 2))
    axis.legend(loc="upper right", frameon=False)
    polish_axes(axis)
    return save_figure(fig, output_dir, "图03_日前购电与储能联合调度结果.png", dpi=dpi)


def figure4_typical_dates(output_dir: Path, dpi: int) -> Path:
    fig, axes = plt.subplots(2, 2, figsize=(12, 8), dpi=dpi, sharex=True, sharey=True)
    for axis, (date_text, title) in zip(axes.ravel(), TARGET_DATES):
        load, pv = _read_attachment2_day(date_text)
        x = _interval_grid()
        axis.plot(x, load, color=COLORS["load"], lw=1.6, label="小区负荷")
        axis.plot(x, pv, color="#2e7d32", lw=1.6, label="光伏实际功率")
        axis.set_title(title)
        axis.set_xlim(0, 24)
        axis.set_xticks(range(0, 25, 4))
        axis.grid(True, alpha=0.3)
        axis.legend(loc="upper right", fontsize=8)
    axes[0, 0].set_ylabel("功率（千瓦）")
    axes[1, 0].set_ylabel("功率（千瓦）")
    axes[1, 0].set_xlabel("时刻（小时）")
    axes[1, 1].set_xlabel("时刻（小时）")
    fig.suptitle("附件2 指定日期的小区负荷与光伏实际功率", y=1.01)
    fig.tight_layout()
    return save_figure(fig, output_dir, "图04_典型日期小区负载与光伏实际功率.png", dpi=dpi)


def _strategy_label(value: str) -> str:
    text = str(value)
    if text in {"版本 A", "A", "完全信息基准（A）"}:
        return "版本 A\n完全信息"
    if "Yesterday-MILP+70" in text:
        return "Yesterday\n+70"
    if "Yesterday-MILP" in text:
        return "Yesterday\n-MILP"
    if "Dynamic" in text:
        return "Dynamic\n（主结果）"
    return text


def figure5_strategy_compare(output_dir: Path, dpi: int) -> Path:
    frame = pd.read_csv(_find_p2_compare())
    required = {"model", "total_cost_yuan", "emergency_purchase_kwh"}
    if not required.issubset(frame.columns):
        raise ValueError(f"p2_compare.csv missing columns: {sorted(required - set(frame.columns))}")
    labels = [_strategy_label(value) for value in frame["model"]]
    total = frame["total_cost_yuan"].to_numpy(float) / 1e4
    emergency = frame["emergency_purchase_kwh"].to_numpy(float) / 1e4
    x = np.arange(len(frame))
    width = 0.34
    fig, axis = plt.subplots(figsize=(10.2, 5.4), constrained_layout=True)
    bars_cost = axis.bar(x - width / 2, total, width, color=COLORS["cost"], label="实际总费用（万元）", zorder=3)
    axis.set_ylabel("实际总费用（万元）", color=COLORS["cost"])
    axis.tick_params(axis="y", labelcolor=COLORS["cost"])
    axis.set_xticks(x)
    axis.set_xticklabels(labels)
    axis.set_ylim(0, max(total) * 1.20)
    polish_axes(axis)
    emergency_axis = axis.twinx()
    bars_emergency = emergency_axis.bar(
        x + width / 2,
        emergency,
        width,
        color="#ef5350",
        label="紧急购电量（万 kWh）",
        zorder=3,
    )
    emergency_axis.set_ylabel("紧急购电量（万 kWh）", color=COLORS["emergency"])
    emergency_axis.tick_params(axis="y", labelcolor=COLORS["emergency"])
    emergency_axis.set_ylim(0, max(float(emergency.max()) * 1.28, 1.0))
    emergency_axis.spines["top"].set_visible(False)
    for bar, value in zip(bars_cost, total):
        axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + max(total) * 0.025, f"{value:.1f}", ha="center", va="bottom", fontsize=9, color=COLORS["cost"])
    for bar, value in zip(bars_emergency, emergency):
        emergency_axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + max(float(emergency.max()) * 0.025, 0.5), f"{value:.1f}" if value else "0", ha="center", va="bottom", fontsize=9, color=COLORS["emergency"])
    axis.legend([bars_cost, bars_emergency], [bars_cost.get_label(), bars_emergency.get_label()], loc="upper right", frameon=False)
    return save_figure(fig, output_dir, "图05_不同购电策略间的总费用与紧急购电量对比.png", dpi=dpi)


def _forecast_grid(values: np.ndarray, release_hour: float) -> np.ndarray:
    result = np.full(INTERVALS, np.nan, dtype=float)
    for offset, value in enumerate(values):
        left = release_hour + offset
        start = max(int(round(left / INTERVAL_HOURS)), 0)
        end = min(start + int(round(1.0 / INTERVAL_HOURS)), INTERVALS)
        if start < end:
            result[start:end] = value
    return result


def figure6_forecast_batches(output_dir: Path, dpi: int) -> Path:
    _, actual = _read_attachment2_day("2025-03-20")
    forecasts = _read_attachment3_day("2025-03-20")
    grid = _interval_grid()
    fig, axes = plt.subplots(2, 2, figsize=(11.2, 7.0), sharex=True, sharey=True, constrained_layout=True)
    axes = axes.ravel()
    actual = np.clip(actual, 0, None)
    maximum = float(actual.max())
    for axis, release_hour in zip(axes, (0.0, 6.0, 12.0, 18.0)):
        forecast = _forecast_grid(forecasts[release_hour], release_hour)
        finite = np.isfinite(forecast)
        if np.any(finite):
            maximum = max(maximum, float(np.nanmax(forecast)))
        axis.plot(grid, actual, color="#1f2937", lw=2.0, label="实际光伏", zorder=3)
        axis.step(grid, forecast, where="post", color="#2563eb", lw=1.8, label="该次发布预报", zorder=4)
        if np.any(finite):
            axis.fill_between(grid, 0, np.where(finite, forecast, np.nan), step="post", color="#93c5fd", alpha=0.35, zorder=2)
        axis.axvline(release_hour, color="#dc2626", ls="--", lw=1.1, alpha=0.75)
        if release_hour > 0:
            axis.axvspan(0, release_hour, color="#f3f4f6", alpha=0.9, zorder=0)
        axis.set_xlim(0, 24)
        axis.set_xticks([0, 4, 8, 12, 16, 20, 24])
        axis.text(0.02, 0.95, f"{int(release_hour):02d}:00 发布", transform=axis.transAxes, ha="left", va="top", fontsize=11, color=COLORS["text"], fontweight="bold")
        polish_axes(axis)
    for axis in axes:
        axis.set_ylim(0, maximum * 1.08)
    axes[0].set_ylabel("光伏功率 / kW")
    axes[2].set_ylabel("光伏功率 / kW")
    axes[2].set_xlabel("日内时间 / h")
    axes[3].set_xlabel("日内时间 / h")
    handles, labels = axes[0].get_legend_handles_labels()
    handles = handles[:2] + [Line2D([0], [0], color="#dc2626", ls="--", lw=1.2)]
    labels = labels[:2] + ["发布时间"]
    fig.legend(handles, labels, loc="upper center", ncol=3, frameon=False, bbox_to_anchor=(0.5, 1.02), fontsize=10)
    return save_figure(fig, output_dir, "图06_典型日光伏预测逐批更新.png", dpi=dpi)


def generate_all(output_dir: Path, *, dpi: int = 220) -> list[Path]:
    """Generate every figure referenced by the final paper."""

    setup_style()
    return [
        figure1_flowchart(output_dir, dpi),
        figure2_typical_day(output_dir, dpi),
        figure3_problem1_dispatch(output_dir, dpi),
        figure4_typical_dates(output_dir, dpi),
        figure5_strategy_compare(output_dir, dpi),
        figure6_forecast_batches(output_dir, dpi),
    ]


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--dpi", type=int, default=220)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.dpi <= 0:
        raise ValueError("--dpi must be positive")
    generate_all(protected_output_path(args.output_dir), dpi=args.dpi)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
