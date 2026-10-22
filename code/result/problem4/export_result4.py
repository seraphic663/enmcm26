"""Excel export for result4-2 / result4-3 into results/problem4/formal."""

from __future__ import annotations

import shutil
from pathlib import Path

from ...common.intervals import date_rows, write_storage_sheet, write_emergency_sheet

import pandas as pd
from openpyxl import load_workbook

from .p4_common import (
    ATTACHMENT_DIR,
    FORMAL_DIR,
    N,
    OUTPUT_START,
)

EPS_E = 1e-7



def export_result4_2(arrays: dict, out_path: Path | None = None) -> Path:
    template = ATTACHMENT_DIR / "附件5" / "result4-2.xlsx"
    out_path = out_path or (FORMAL_DIR / "result4-2.xlsx")
    dates = [pd.Timestamp(str(d)) for d in arrays["dates"]]
    out_idx = [i for i, d in enumerate(dates) if d >= OUTPUT_START]
    if len(out_idx) != 334:
        raise ValueError(f"输出天数应为334，实际{len(out_idx)}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out_path)
    wb = load_workbook(out_path)

    ws = wb["计划购电量"]
    date_row = date_rows(ws)
    for i in out_idx:
        r = date_row[dates[i].normalize()]
        x = arrays["x"][i]
        for t in range(N):
            ws.cell(r, t + 2, float(x[t]))
        ws.cell(r, N + 2, float(x.sum()))
        ws.cell(r, N + 3, float(arrays["planned_cost"][i]))

    write_storage_sheet(
        wb, dates, out_idx, arrays["charge"], arrays["discharge"], arrays["soc"]
    )
    write_emergency_sheet(wb, dates, out_idx, arrays["emergency"], 2, EPS_E)
    wb.save(out_path)
    return out_path


def export_result4_3(arrays: dict, out_path: Path | None = None) -> Path:
    template = ATTACHMENT_DIR / "附件5" / "result4-3.xlsx"
    out_path = out_path or (FORMAL_DIR / "result4-3.xlsx")
    dates = [pd.Timestamp(str(d)) for d in arrays["dates"]]
    out_idx = [i for i, d in enumerate(dates) if d >= OUTPUT_START]
    if len(out_idx) != 334:
        raise ValueError(f"输出天数应为334，实际{len(out_idx)}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out_path)
    wb = load_workbook(out_path)

    ws = wb["计划购电量"]
    date_row = date_rows(ws)
    for i in out_idx:
        r = date_row[dates[i].normalize()]
        xp = arrays["x_plan"][i]
        for t in range(N):
            ws.cell(r, t + 2, float(xp[t]))
        ws.cell(r, N + 2, float(xp.sum()))
        ws.cell(r, N + 3, float(arrays["plan_ref_cost"][i]))

    ws = wb["调整购电量"]
    date_row = date_rows(ws)
    for i in out_idx:
        r = date_row[dates[i].normalize()]
        xa = arrays["x_act"][i]
        for t in range(N):
            ws.cell(r, t + 2, float(xa[t]))
        ws.cell(r, N + 2, float(xa.sum()))
        ws.cell(r, N + 3, float(arrays["cost_total"][i]))

    write_storage_sheet(
        wb, dates, out_idx, arrays["charge"], arrays["discharge"], arrays["soc"]
    )
    write_emergency_sheet(wb, dates, out_idx, arrays["emergency"], 3, EPS_E)
    wb.save(out_path)
    return out_path
