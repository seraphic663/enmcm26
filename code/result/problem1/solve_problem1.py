"""Solve and audit CUMCM 2026 Problem C, Question 1.

P0 rule: audit the time convention before building either MILP.  The official
materials are internally inconsistent, so the script supports two explicit
modes and never changes the official template's time labels:

* A: each attachment timestamp is the right endpoint of its 10-minute period;
* B: each attachment timestamp is the left endpoint of its 10-minute period.

Every run solves the shared numeric model once and compares both time mappings.  By default it exports both
``result-A.xlsx`` and ``result-B.xlsx``; use ``--time-mapping A`` or
``--time-mapping B`` to export only one candidate.
"""

from __future__ import annotations

import argparse
import shutil
from dataclasses import dataclass, replace
from datetime import time
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
import pulp
from openpyxl import load_workbook
from scipy.optimize import linprog

from ...common.paths import project_root, results_root
from ...common.dataset import sha256


INTERVAL_COUNT = 144
DELTA_HOURS = 1.0 / 6.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
SOC_MIN_KWH = 1200.0
SOC_MAX_KWH = 10800.0
SOC_INITIAL_KWH = 6000.0
POWER_MAX_KW = 5000.0
ENERGY_MAX_KWH = POWER_MAX_KW * DELTA_HOURS
MODEL_TOL = 5e-5
EXPORT_TOL = 1e-8

EXPECTED_COLUMNS = ["时间", "电价", "小区负载", "光伏发电预测功率"]
PURCHASE_SHEET = "计划购电量"
STORAGE_SHEET = "充放电量"
TABLE1_INTERVALS = [
    "10:00-10:10",
    "12:00-12:10",
    "14:00-14:10",
    "16:00-16:10",
    "18:00-18:10",
    "20:00-20:10",
]
OFFICIAL_FOUR_HOUR_PERIODS = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]


@dataclass(frozen=True)
class Paths:
    problem_dir: Path
    input_workbook: Path
    template_dir: Path
    template_workbook: Path
    output_dir: Path
    formal_dir: Path
    audit_dir: Path


@dataclass(frozen=True)
class MappingSpec:
    code: str
    description: str
    source_label_role: str
    start_offset_minutes: int

    @property
    def boundary_start(self) -> str:
        return clock_label(self.start_offset_minutes)

    @property
    def boundary_end(self) -> str:
        if self.code == "A":
            return "24:00"
        return clock_label(self.start_offset_minutes + 24 * 60)

    @property
    def interval_labels(self) -> list[str]:
        return [
            interval_label(self.start_offset_minutes + index * 10)
            for index in range(INTERVAL_COUNT)
        ]

    @property
    def block_labels(self) -> list[str]:
        if self.code == "A":
            return OFFICIAL_FOUR_HOUR_PERIODS.copy()
        return [
            f"{clock_label(self.start_offset_minutes + block * 240)}-"
            f"{clock_label(self.start_offset_minutes + (block + 1) * 240)}"
            for block in range(6)
        ]


MAPPING_A = MappingSpec(
    code="A",
    description="附件时间为10分钟区间右端点",
    source_label_role="right_endpoint",
    start_offset_minutes=0,
)
MAPPING_B = MappingSpec(
    code="B",
    description="附件时间为10分钟区间左端点",
    source_label_role="left_endpoint",
    start_offset_minutes=10,
)
MAPPINGS = {spec.code: spec for spec in (MAPPING_A, MAPPING_B)}


@dataclass(frozen=True)
class InputData:
    source_labels: list[str]
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray


@dataclass(frozen=True)
class Solution:
    mapping: MappingSpec
    status: str
    purchase_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    curtailed_kwh: np.ndarray
    charge_state: np.ndarray
    soc_kwh: np.ndarray
    objective_cost: float


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--time-mapping",
        choices=("A", "B", "both"),
        default="both",
        help="选择导出口径；默认both，同时生成result-A.xlsx和result-B.xlsx",
    )
    return parser.parse_args()


def get_paths() -> Paths:
    problem_dir = project_root()
    template_dir = problem_dir / "附件" / "附件5"
    return Paths(
        problem_dir=problem_dir,
        input_workbook=problem_dir / "附件" / "附件1.xlsx",
        template_dir=template_dir,
        template_workbook=template_dir / "result1.xlsx",
        output_dir=results_root() / "problem1",
        formal_dir=results_root() / "problem1" / "formal",
        audit_dir=results_root() / "problem1" / "audit",
    )


def mode_output_paths(paths: Paths, mapping: MappingSpec) -> tuple[Path, Path]:
    return (
        paths.formal_dir / f"result-{mapping.code}.xlsx",
        paths.audit_dir / f"problem1_schedule-{mapping.code}.csv",
    )


def clock_label(total_minutes: int) -> str:
    day, within_day = divmod(total_minutes, 24 * 60)
    hour, minute = divmod(within_day, 60)
    suffix = f"+{day}" if day else ""
    return f"{hour}:{minute:02d}{suffix}"


def interval_label(start_minutes: int) -> str:
    return f"{clock_label(start_minutes)}-{clock_label(start_minutes + 10)}"


def parse_source_minutes(value: Any) -> int:
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    if text == "0:00+1":
        return 24 * 60
    if text.endswith(":00") and text.count(":") == 2:
        text = text[:-3]
    parts = text.split(":")
    if len(parts) != 2:
        raise AssertionError(f"无法识别时间标签: {value!r}")
    hour, minute = (int(part) for part in parts)
    assert 0 <= hour <= 23 and 0 <= minute <= 59, f"非法时间标签: {value!r}"
    return hour * 60 + minute


def read_and_validate_input(path: Path) -> InputData:
    frame = pd.read_excel(path, sheet_name="Sheet1")
    assert list(frame.columns) == EXPECTED_COLUMNS, f"附件1列名不符: {list(frame.columns)!r}"
    assert len(frame) == INTERVAL_COUNT, f"附件1应有144行，实际{len(frame)}行"
    assert not frame.isna().any().any(), "附件1存在空值"

    source_labels: list[str] = []
    for index, raw in enumerate(frame["时间"].tolist()):
        actual = parse_source_minutes(raw)
        expected = (index + 1) * 10
        assert actual == expected, f"附件1第{index + 1}行时间错位: {raw!r}"
        source_labels.append(clock_label(actual))
    assert source_labels[0] == "0:10" and source_labels[-1] == "0:00+1"

    numeric = frame[EXPECTED_COLUMNS[1:]].apply(pd.to_numeric, errors="raise")
    assert np.isfinite(numeric.to_numpy(dtype=float)).all(), "附件1含非有限数值"
    price = numeric["电价"].to_numpy(dtype=float)
    load_kw = numeric["小区负载"].to_numpy(dtype=float)
    pv_kw = numeric["光伏发电预测功率"].to_numpy(dtype=float)
    assert np.all(price >= 0) and np.all(load_kw >= 0) and np.all(pv_kw >= 0)
    return InputData(
        source_labels=source_labels,
        price=price,
        load_kw=load_kw,
        pv_kw=pv_kw,
        load_kwh=load_kw * DELTA_HOURS,
        pv_kwh=pv_kw * DELTA_HOURS,
    )


def read_row_values(worksheet: Any, row: int, start: int, count: int) -> list[Any]:
    return [worksheet.cell(row=row, column=start + offset).value for offset in range(count)]


def audit_time_conventions(paths: Paths) -> tuple[InputData, dict[str, Any]]:
    """P0 audit: this must finish before solve_milp is called."""
    required = [
        paths.input_workbook,
        paths.template_workbook,
    ]
    required.extend(paths.problem_dir.joinpath("附件").glob("附件[234].xlsx"))
    required.extend(paths.template_dir.glob("result*.xlsx"))
    missing = [str(path) for path in required if not path.is_file()]
    assert not missing, f"P0审计缺少文件: {missing}"

    data = read_and_validate_input(paths.input_workbook)

    endpoint_workbooks: dict[str, list[str]] = {"附件1.xlsx": data.source_labels}
    for filename in ("附件2.xlsx", "附件4.xlsx"):
        path = paths.problem_dir / "附件" / filename
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                raw_headers = read_row_values(worksheet, 1, 2, INTERVAL_COUNT)
                parsed = [clock_label(parse_source_minutes(value)) for value in raw_headers]
                assert parsed == data.source_labels, f"{filename}/{worksheet.title}时间头不一致"
            endpoint_workbooks[filename] = data.source_labels
        finally:
            workbook.close()

    b_labels = MAPPING_B.interval_labels
    template_plan_findings: dict[str, dict[str, Any]] = {}
    storage_findings: dict[str, dict[str, Any]] = {}
    template_files = [
        "result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"
    ]
    for filename in template_files:
        path = paths.template_dir / filename
        workbook = load_workbook(path, read_only=False, data_only=False)
        try:
            plan_sheets = [
                name for name in ("计划购电量", "调整购电量") if name in workbook.sheetnames
            ]
            plan_details = []
            for sheet_name in plan_sheets:
                worksheet = workbook[sheet_name]
                if filename == "result1.xlsx":
                    labels = [worksheet.cell(row=row, column=1).value for row in range(2, 146)]
                else:
                    labels = read_row_values(worksheet, 1, 2, INTERVAL_COUNT)
                exact_b_count = sum(a == b for a, b in zip(labels, b_labels, strict=True))
                differences = [
                    {
                        "one_based_position": index + 1,
                        "actual": actual,
                        "expected_B": expected,
                    }
                    for index, (actual, expected) in enumerate(
                        zip(labels, b_labels, strict=True)
                    )
                    if actual != expected
                ]
                plan_details.append(
                    {
                        "sheet": sheet_name,
                        "count": len(labels),
                        "first": labels[0],
                        "penultimate": labels[-2],
                        "last": labels[-1],
                        "exact_B_label_count": exact_b_count,
                        "differences": differences,
                    }
                )
                assert len(labels) == INTERVAL_COUNT
                assert labels[0] == "0:10-0:20"
                assert labels[-2] == "23:50-0:00+1"
                if filename == "result1.xlsx":
                    assert labels == b_labels
                else:
                    assert differences == [
                        {
                            "one_based_position": 42,
                            "actual": "7:0-7:10",
                            "expected_B": "7:00-7:10",
                        },
                        {
                            "one_based_position": 144,
                            "actual": "0:00-0:10+1",
                            "expected_B": "0:00+1-0:10+1",
                        },
                    ]
            template_plan_findings[filename] = {"sheets": plan_details}

            storage = workbook[STORAGE_SHEET]
            values = [
                str(cell.value).strip()
                for row in storage.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            for label in OFFICIAL_FOUR_HOUR_PERIODS:
                assert label in values, f"{filename}缺少储能汇总标签{label}"
            assert "24:00" in values
            assert any(value in ("0:00", "00:00:00") for value in values)
            storage_findings[filename] = {
                "civil_day_blocks": True,
                "has_0000": True,
                "has_2400": True,
            }
        finally:
            workbook.close()

    audit = {
        "completed_before_model": True,
        "attachment_endpoint_workbooks": endpoint_workbooks,
        "template_plan_findings": template_plan_findings,
        "storage_findings": storage_findings,
        "statement_requires_civil_day_soc": True,
        "statement_defines_endpoint_role": False,
        "all_plan_templates_start_at_0010": True,
        "all_storage_templates_use_0000_2400": True,
        "result234_header_anomalies": ["7:0-7:10", "0:00-0:10+1"],
        "unique_interpretation": False,
        "conclusion": (
            "A与题目全天/0:00-24:00及表2一致；B与全部计划购电模板逐行约定一致。"
            "题面未声明时间戳端点角色，无法唯一判定。"
        ),
    }
    return data, audit


def print_p0_audit(audit: dict[str, Any]) -> None:
    print("=== P0 时间索引审计（先于MILP） ===")
    print("附件1/2/4: 均为144个时间标签，0:10起、0:00+1止")
    print("result1/2/3/4: 所有计划/调整购电模板均从0:10-0:20起")
    print("result1最后一段: 0:00+1-0:10+1")
    print("result2/3/4另有两处统一拼写异常: 7:0-7:10；0:00-0:10+1")
    print("全部充放电模板: 明确使用0:00-4:00...20:00-24:00及0:00/24:00 SOC")
    print(f"P0结论: {audit['conclusion']}")
    print("以下使用同一数值解核对A（右端点）与B（左端点）标签，不改官方模板。\n")


def solve_highs_lp_relaxation(
    data: InputData,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, float]:
    """Solve a high-precision LP lower bound and prove it is MILP-feasible.

    The binary charge/discharge state is omitted in this independent model.
    If the returned optimum has no simultaneous charge and discharge, assigning
    the state from the active flow makes it feasible for the original MILP.
    Hence the LP lower bound is also a MILP upper bound and certifies optimality.
    """
    count = INTERVAL_COUNT
    purchase_offset = 0
    charge_offset = count
    discharge_offset = 2 * count
    curtail_offset = 3 * count
    soc_offset = 4 * count
    variable_count = 4 * count + (count + 1)

    objective = np.zeros(variable_count)
    objective[purchase_offset:purchase_offset + count] = data.price
    equality_rows: list[np.ndarray] = []
    equality_values: list[float] = []
    for t in range(count):
        balance = np.zeros(variable_count)
        balance[purchase_offset + t] = 1.0
        balance[charge_offset + t] = -1.0
        balance[discharge_offset + t] = 1.0
        balance[curtail_offset + t] = -1.0
        equality_rows.append(balance)
        equality_values.append(float(data.load_kwh[t] - data.pv_kwh[t]))

        update = np.zeros(variable_count)
        update[soc_offset + t + 1] = 1.0
        update[soc_offset + t] = -1.0
        update[charge_offset + t] = -ETA_CHARGE
        update[discharge_offset + t] = 1.0 / ETA_DISCHARGE
        equality_rows.append(update)
        equality_values.append(0.0)

    initial = np.zeros(variable_count)
    initial[soc_offset] = 1.0
    equality_rows.append(initial)
    equality_values.append(SOC_INITIAL_KWH)
    terminal = np.zeros(variable_count)
    terminal[soc_offset + count] = 1.0
    equality_rows.append(terminal)
    equality_values.append(SOC_INITIAL_KWH)

    bounds: list[tuple[float, float | None]] = []
    bounds.extend([(0.0, None)] * count)
    bounds.extend([(0.0, ENERGY_MAX_KWH)] * count)
    bounds.extend([(0.0, ENERGY_MAX_KWH)] * count)
    bounds.extend((0.0, float(value)) for value in data.pv_kwh)
    bounds.extend([(SOC_MIN_KWH, SOC_MAX_KWH)] * (count + 1))

    result = linprog(
        objective,
        A_eq=np.array(equality_rows),
        b_eq=np.array(equality_values),
        bounds=bounds,
        method="highs",
        options={
            "primal_feasibility_tolerance": 1e-9,
            "dual_feasibility_tolerance": 1e-9,
            "ipm_optimality_tolerance": 1e-10,
        },
    )
    assert result.success, f"HiGHS LP独立校验失败: {result.message}"
    vector = np.asarray(result.x, dtype=float)
    purchase = vector[purchase_offset:purchase_offset + count]
    charge = vector[charge_offset:charge_offset + count]
    discharge = vector[discharge_offset:discharge_offset + count]
    curtailed = vector[curtail_offset:curtail_offset + count]
    soc = vector[soc_offset:soc_offset + count + 1]
    simultaneous = (charge > EXPORT_TOL) & (discharge > EXPORT_TOL)
    assert not simultaneous.any(), "HiGHS LP下界解存在同时充放电，不能证明MILP最优"
    state = (charge > EXPORT_TOL).astype(float)
    return purchase, charge, discharge, curtailed, state, soc, float(result.fun)


def solve_milp(data: InputData, mapping: MappingSpec) -> Solution:
    model = pulp.LpProblem(f"CUMCM_2026_C_Problem1_{mapping.code}", pulp.LpMinimize)
    periods = range(INTERVAL_COUNT)
    purchase = pulp.LpVariable.dicts("purchase_kwh", periods, lowBound=0.0)
    charge = pulp.LpVariable.dicts("charge_kwh", periods, lowBound=0.0)
    discharge = pulp.LpVariable.dicts("discharge_kwh", periods, lowBound=0.0)
    curtailed = pulp.LpVariable.dicts("curtailed_kwh", periods, lowBound=0.0)
    charge_state = pulp.LpVariable.dicts("charge_state", periods, cat="Binary")
    soc = pulp.LpVariable.dicts(
        "soc_kwh", range(INTERVAL_COUNT + 1), lowBound=SOC_MIN_KWH, upBound=SOC_MAX_KWH
    )

    model += pulp.lpSum(data.price[t] * purchase[t] for t in periods)
    model += soc[0] == SOC_INITIAL_KWH, "initial_soc"
    for t in periods:
        model += (
            purchase[t] + data.pv_kwh[t] + discharge[t]
            == data.load_kwh[t] + charge[t] + curtailed[t]
        ), f"energy_balance_{t:03d}"
        model += (
            soc[t + 1] == soc[t] + ETA_CHARGE * charge[t]
            - discharge[t] / ETA_DISCHARGE
        ), f"soc_update_{t:03d}"
        model += charge[t] <= ENERGY_MAX_KWH * charge_state[t]
        model += discharge[t] <= ENERGY_MAX_KWH * (1 - charge_state[t])
        model += curtailed[t] <= data.pv_kwh[t]
    model += soc[INTERVAL_COUNT] == SOC_INITIAL_KWH, "terminal_soc"

    status_code = model.solve(
        pulp.PULP_CBC_CMD(msg=False, threads=1, gapRel=0.0, options=["randomSeed 0"])
    )
    status = pulp.LpStatus[status_code]
    assert status == "Optimal", f"模式{mapping.code}求解状态为{status}"

    raw_objective = float(pulp.value(model.objective))
    (
        precise_purchase,
        precise_charge,
        precise_discharge,
        precise_curtailment,
        precise_state,
        precise_soc,
        precise_objective,
    ) = solve_highs_lp_relaxation(data)
    assert abs(precise_objective - raw_objective) <= 0.01, (
        "HiGHS高精度目标偏离CBC最优目标超过0.01元: "
        f"{precise_objective - raw_objective:.12g}"
    )

    return Solution(
        mapping=mapping,
        status=status,
        purchase_kwh=precise_purchase,
        charge_kwh=precise_charge,
        discharge_kwh=precise_discharge,
        curtailed_kwh=precise_curtailment,
        charge_state=precise_state,
        soc_kwh=precise_soc,
        objective_cost=precise_objective,
    )


def assert_max_abs(name: str, values: np.ndarray, tolerance: float) -> float:
    maximum = float(np.max(np.abs(values))) if len(values) else 0.0
    assert maximum <= tolerance, f"{name}最大绝对误差{maximum:.12g}超过{tolerance:g}"
    return maximum


def validate_numeric_solution(
    data: InputData,
    purchase: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    curtailed: np.ndarray,
    state: np.ndarray,
    soc_start: np.ndarray,
    soc_end: np.ndarray,
    tolerance: float,
) -> dict[str, float | bool]:
    arrays = (purchase, charge, discharge, curtailed, state, soc_start, soc_end)
    assert all(len(array) == INTERVAL_COUNT for array in arrays)
    assert all(np.isfinite(array).all() for array in arrays)
    assert min(np.min(purchase), np.min(charge), np.min(discharge), np.min(curtailed)) >= -tolerance
    assert np.max(curtailed - data.pv_kwh) <= tolerance
    assert np.max(charge) <= ENERGY_MAX_KWH + tolerance
    assert np.max(discharge) <= ENERGY_MAX_KWH + tolerance
    simultaneous = (charge > tolerance) & (discharge > tolerance)
    assert not simultaneous.any(), "存在同时充放电"
    assert np.max(np.minimum(np.abs(state), np.abs(state - 1))) <= tolerance
    assert np.max(charge - ENERGY_MAX_KWH * state) <= tolerance
    assert np.max(discharge - ENERGY_MAX_KWH * (1 - state)) <= tolerance
    assert abs(soc_start[0] - SOC_INITIAL_KWH) <= tolerance
    assert abs(soc_end[-1] - SOC_INITIAL_KWH) <= tolerance
    assert min(np.min(soc_start), np.min(soc_end)) >= SOC_MIN_KWH - tolerance
    assert max(np.max(soc_start), np.max(soc_end)) <= SOC_MAX_KWH + tolerance

    continuity = assert_max_abs("SOC连续性", soc_start[1:] - soc_end[:-1], tolerance)
    soc_residual = assert_max_abs(
        "SOC递推",
        soc_end - soc_start - ETA_CHARGE * charge + discharge / ETA_DISCHARGE,
        tolerance,
    )
    balance = assert_max_abs(
        "电量平衡",
        purchase + data.pv_kwh + discharge - data.load_kwh - charge - curtailed,
        tolerance,
    )
    return {
        "balance_max_abs_kwh": balance,
        "soc_update_max_abs_kwh": soc_residual,
        "soc_continuity_max_abs_kwh": continuity,
        "simultaneous_charge_discharge": bool(simultaneous.any()),
    }


def summarize(data: InputData, solution: Solution) -> dict[str, Any]:
    labels = solution.mapping.interval_labels
    index_by_label = {label: index for index, label in enumerate(labels)}
    assert all(label in index_by_label for label in TABLE1_INTERVALS)
    table1 = {label: float(solution.purchase_kwh[index_by_label[label]]) for label in TABLE1_INTERVALS}
    expected_indices = [60, 72, 84, 96, 108, 120] if solution.mapping.code == "A" else [59, 71, 83, 95, 107, 119]
    assert [index_by_label[label] for label in TABLE1_INTERVALS] == expected_indices

    positional_blocks: dict[str, dict[str, float]] = {}
    for block, label in enumerate(solution.mapping.block_labels):
        start, stop = block * 24, (block + 1) * 24
        positional_blocks[label] = {
            "charge_kwh": float(np.sum(solution.charge_kwh[start:stop])),
            "discharge_kwh": float(np.sum(solution.discharge_kwh[start:stop])),
        }
    all_soc = solution.soc_kwh
    return {
        "total_purchase_kwh": float(np.sum(solution.purchase_kwh)),
        "total_cost_yuan": float(np.dot(data.price, solution.purchase_kwh)),
        "table1": table1,
        "table1_indices_zero_based": expected_indices,
        "positional_blocks": positional_blocks,
        "official_block_labels_aligned": solution.mapping.block_labels == OFFICIAL_FOUR_HOUR_PERIODS,
        "purchase_template_labels_aligned": solution.mapping.interval_labels == MAPPING_B.interval_labels,
        "soc_start_label": solution.mapping.boundary_start,
        "soc_end_label": solution.mapping.boundary_end,
        "soc_start_kwh": float(all_soc[0]),
        "soc_end_kwh": float(all_soc[-1]),
        "soc_min_kwh": float(np.min(all_soc)),
        "soc_max_kwh": float(np.max(all_soc)),
        "max_charge_kwh": float(np.max(solution.charge_kwh)),
        "max_discharge_kwh": float(np.max(solution.discharge_kwh)),
    }


def solution_frame(data: InputData, solution: Solution) -> pd.DataFrame:
    labels = solution.mapping.interval_labels
    return pd.DataFrame(
        {
            "time_mapping": solution.mapping.code,
            "source_label_role": solution.mapping.source_label_role,
            "period_index": np.arange(1, INTERVAL_COUNT + 1),
            "source_time_label": data.source_labels,
            "interpreted_interval": labels,
            "price_yuan_per_kwh": data.price,
            "load_kw": data.load_kw,
            "pv_kw": data.pv_kw,
            "load_kwh": data.load_kwh,
            "pv_kwh": data.pv_kwh,
            "purchase_kwh": solution.purchase_kwh,
            "charge_kwh": solution.charge_kwh,
            "discharge_kwh": solution.discharge_kwh,
            "curtailed_kwh": solution.curtailed_kwh,
            "charge_state": solution.charge_state,
            "soc_start_label": [clock_label(solution.mapping.start_offset_minutes + i * 10) for i in range(INTERVAL_COUNT)],
            "soc_end_label": [clock_label(solution.mapping.start_offset_minutes + (i + 1) * 10) for i in range(INTERVAL_COUNT)],
            "soc_start_kwh": solution.soc_kwh[:-1],
            "soc_end_kwh": solution.soc_kwh[1:],
        }
    )


def workbook_layout_snapshot(workbook: openpyxl.Workbook) -> dict[str, Any]:
    result: dict[str, Any] = {
        "sheetnames": tuple(workbook.sheetnames),
        "active": workbook.index(workbook.active),
        "defined_names": tuple(sorted(str(item) for item in workbook.defined_names.values())),
        "sheets": {},
    }
    for worksheet in workbook.worksheets:
        styles = tuple(
            # An unstyled blank cell has _style=None in the source.  After a
            # value is written and the workbook is re-opened, openpyxl exposes
            # the visually identical all-zero default StyleArray.  Canonicalize
            # both as unstyled while still comparing every non-default style.
            (
                cell.coordinate,
                repr(cell._style) if cell.has_style else None,
                cell.number_format,
            )
            for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column)
            for cell in row
        )
        row_dims = tuple(sorted((i, d.height, d.hidden, d.outlineLevel, d.collapsed) for i, d in worksheet.row_dimensions.items()))
        col_dims = tuple(sorted((i, d.width, d.hidden, d.outlineLevel, d.collapsed) for i, d in worksheet.column_dimensions.items()))
        result["sheets"][worksheet.title] = {
            "dimensions": (worksheet.max_row, worksheet.max_column),
            "state": worksheet.sheet_state,
            "merges": tuple(str(item) for item in worksheet.merged_cells.ranges),
            "freeze": str(worksheet.freeze_panes),
            "filter": worksheet.auto_filter.ref,
            "print_area": str(worksheet.print_area),
            "print_titles": (worksheet.print_title_rows, worksheet.print_title_cols),
            "row_dims": row_dims,
            "col_dims": col_dims,
            "sheet_format": repr(worksheet.sheet_format),
            "sheet_properties": repr(worksheet.sheet_properties),
            "sheet_view": repr(worksheet.sheet_view),
            "page_margins": repr(worksheet.page_margins),
            "page_setup": repr(worksheet.page_setup),
            "print_options": repr(worksheet.print_options),
            "protection": repr(worksheet.protection),
            "styles": styles,
        }
    return result


def workbook_values(workbook: openpyxl.Workbook) -> dict[tuple[str, str], Any]:
    return {
        (worksheet.title, cell.coordinate): cell.value
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column)
        for cell in row
    }


def template_baseline(path: Path) -> tuple[dict[str, Any], dict[tuple[str, str], Any], str, list[str]]:
    workbook = load_workbook(path, data_only=False)
    try:
        labels = [workbook[PURCHASE_SHEET].cell(row=row, column=1).value for row in range(2, 146)]
        return workbook_layout_snapshot(workbook), workbook_values(workbook), sha256(path), labels
    finally:
        workbook.close()


def write_candidate_workbook(
    paths: Paths,
    solution: Solution,
    summary: dict[str, Any],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(paths.template_workbook, output_path)
    workbook = load_workbook(output_path)
    try:
        assert workbook.sheetnames == [PURCHASE_SHEET, STORAGE_SHEET]
        purchase_sheet = workbook[PURCHASE_SHEET]
        storage_sheet = workbook[STORAGE_SHEET]
        assert [purchase_sheet.cell(row=row, column=1).value for row in range(2, 146)] == MAPPING_B.interval_labels
        for index, value in enumerate(solution.purchase_kwh, start=2):
            purchase_sheet.cell(row=index, column=2, value=float(value))

        for block, model_label in enumerate(solution.mapping.block_labels):
            row = block + 2
            values = summary["positional_blocks"][model_label]
            storage_sheet.cell(row=row, column=2, value=values["charge_kwh"])
            storage_sheet.cell(row=row, column=3, value=values["discharge_kwh"])
        storage_sheet["E2"] = summary["soc_start_kwh"]
        storage_sheet["E3"] = summary["soc_end_kwh"]
        workbook.save(output_path)
    finally:
        workbook.close()


def audit_serialized_schedule(
    csv_path: Path,
    input_path: Path,
    mapping: MappingSpec,
    expected_objective: float,
) -> tuple[pd.DataFrame, InputData, dict[str, float | bool]]:
    data = read_and_validate_input(input_path)
    frame = pd.read_csv(csv_path, encoding="utf-8-sig")
    assert len(frame) == INTERVAL_COUNT
    assert frame["time_mapping"].astype(str).unique().tolist() == [mapping.code]
    assert frame["source_label_role"].unique().tolist() == [mapping.source_label_role]
    assert frame["period_index"].tolist() == list(range(1, INTERVAL_COUNT + 1))
    assert frame["source_time_label"].astype(str).tolist() == data.source_labels
    assert frame["interpreted_interval"].astype(str).tolist() == mapping.interval_labels

    numeric = [
        "price_yuan_per_kwh", "load_kw", "pv_kw", "load_kwh", "pv_kwh",
        "purchase_kwh", "charge_kwh", "discharge_kwh", "curtailed_kwh",
        "charge_state", "soc_start_kwh", "soc_end_kwh",
    ]
    for column in numeric:
        assert np.isfinite(frame[column].to_numpy(dtype=float)).all()
    assert_max_abs("CSV电价", frame["price_yuan_per_kwh"].to_numpy() - data.price, EXPORT_TOL)
    assert_max_abs("CSV负载功率", frame["load_kw"].to_numpy() - data.load_kw, EXPORT_TOL)
    assert_max_abs("CSV光伏功率", frame["pv_kw"].to_numpy() - data.pv_kw, EXPORT_TOL)
    assert_max_abs("负载kW转kWh", frame["load_kwh"].to_numpy() - frame["load_kw"].to_numpy() / 6, EXPORT_TOL)
    assert_max_abs("光伏kW转kWh", frame["pv_kwh"].to_numpy() - frame["pv_kw"].to_numpy() / 6, EXPORT_TOL)

    checks = validate_numeric_solution(
        data,
        frame["purchase_kwh"].to_numpy(dtype=float),
        frame["charge_kwh"].to_numpy(dtype=float),
        frame["discharge_kwh"].to_numpy(dtype=float),
        frame["curtailed_kwh"].to_numpy(dtype=float),
        frame["charge_state"].to_numpy(dtype=float),
        frame["soc_start_kwh"].to_numpy(dtype=float),
        frame["soc_end_kwh"].to_numpy(dtype=float),
        MODEL_TOL,
    )
    objective = float(np.dot(data.price, frame["purchase_kwh"].to_numpy(dtype=float)))
    assert abs(objective - expected_objective) <= MODEL_TOL
    return frame, data, checks


def audit_candidate_workbook(
    paths: Paths,
    mapping: MappingSpec,
    output_path: Path,
    serialized: pd.DataFrame,
    data: InputData,
    summary: dict[str, Any],
    baseline_layout: dict[str, Any],
    baseline_values: dict[tuple[str, str], Any],
    template_hash: str,
    template_labels: list[str],
) -> dict[str, Any]:
    assert sha256(paths.template_workbook) == template_hash, "原模板哈希变化"
    workbook = load_workbook(output_path, data_only=False)
    try:
        assert workbook_layout_snapshot(workbook) == baseline_layout, "模板结构或格式被改变"
        allowed = {(PURCHASE_SHEET, f"B{row}") for row in range(2, 146)}
        allowed |= {(STORAGE_SHEET, f"{column}{row}") for column in ("B", "C") for row in range(2, 8)}
        allowed |= {(STORAGE_SHEET, "E2"), (STORAGE_SHEET, "E3")}
        output_values = workbook_values(workbook)
        for key, value in baseline_values.items():
            if key not in allowed:
                assert output_values[key] == value, f"非目标单元格被修改: {key}"

        purchase_sheet = workbook[PURCHASE_SHEET]
        exported_labels = [purchase_sheet.cell(row=row, column=1).value for row in range(2, 146)]
        assert exported_labels == template_labels == MAPPING_B.interval_labels
        exported_purchase = np.array(
            [purchase_sheet.cell(row=row, column=2).value for row in range(2, 146)], dtype=float
        )
        assert_max_abs(
            "Excel购电量与CSV", exported_purchase - serialized["purchase_kwh"].to_numpy(dtype=float), EXPORT_TOL
        )

        storage = workbook[STORAGE_SHEET]
        for block, model_label in enumerate(mapping.block_labels):
            row = block + 2
            values = summary["positional_blocks"][model_label]
            assert abs(float(storage.cell(row=row, column=2).value) - values["charge_kwh"]) <= EXPORT_TOL
            assert abs(float(storage.cell(row=row, column=3).value) - values["discharge_kwh"]) <= EXPORT_TOL
        assert abs(float(storage["E2"].value) - SOC_INITIAL_KWH) <= EXPORT_TOL
        assert abs(float(storage["E3"].value) - SOC_INITIAL_KWH) <= EXPORT_TOL

        error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
        formula_errors = [
            f"{worksheet.title}!{cell.coordinate}"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens)
        ]
        assert not formula_errors

        export_checks = validate_numeric_solution(
            data,
            exported_purchase,
            serialized["charge_kwh"].to_numpy(dtype=float),
            serialized["discharge_kwh"].to_numpy(dtype=float),
            serialized["curtailed_kwh"].to_numpy(dtype=float),
            serialized["charge_state"].to_numpy(dtype=float),
            serialized["soc_start_kwh"].to_numpy(dtype=float),
            serialized["soc_end_kwh"].to_numpy(dtype=float),
            MODEL_TOL,
        )
        return {
            "sheetnames": workbook.sheetnames,
            "template_format_preserved": True,
            "non_target_values_preserved": True,
            "formula_error_count": 0,
            "purchase_template_labels_aligned": exported_labels == mapping.interval_labels,
            "official_storage_labels_aligned": mapping.block_labels == OFFICIAL_FOUR_HOUR_PERIODS,
            **export_checks,
        }
    finally:
        workbook.close()


def write_comparison(
    path: Path,
    solutions: dict[str, Solution],
    summaries: dict[str, dict[str, Any]],
) -> pd.DataFrame:
    assert set(solutions) == {"A", "B"}
    a, b = solutions["A"], solutions["B"]
    sa, sb = summaries["A"], summaries["B"]
    rows: list[dict[str, Any]] = [
        {"metric": "objective_cost_yuan", "A": sa["total_cost_yuan"], "B": sb["total_cost_yuan"]},
        {"metric": "total_purchase_kwh", "A": sa["total_purchase_kwh"], "B": sb["total_purchase_kwh"]},
    ]
    for label in TABLE1_INTERVALS:
        rows.append({"metric": f"table1_purchase_kwh::{label}", "A": sa["table1"][label], "B": sb["table1"][label]})
    rows.extend(
        [
            {
                "metric": "soc_model_index_max_abs_difference_kwh",
                "A": 0.0,
                "B": float(np.max(np.abs(a.soc_kwh - b.soc_kwh))),
            },
            {
                "metric": "soc_same_clock_0010_to_2400_max_abs_difference_kwh",
                "A": 0.0,
                "B": float(np.max(np.abs(a.soc_kwh[1:] - b.soc_kwh[:-1]))),
            },
        ]
    )
    frame = pd.DataFrame(rows)
    frame["B_minus_A"] = frame["B"] - frame["A"]
    frame.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.12f")
    return frame


def print_mode_report(
    solution: Solution,
    summary: dict[str, Any],
    checks: dict[str, float | bool],
    export_audit: dict[str, Any],
    workbook_path: Path,
    csv_path: Path,
) -> None:
    mapping = solution.mapping
    print(f"=== 模式{mapping.code}: {mapping.description} ===")
    print(f"模型时间范围: {mapping.boundary_start} 至 {mapping.boundary_end}")
    print(f"求解状态: {solution.status}")
    print(f"全天/模型周期购电量: {summary['total_purchase_kwh']:.6f} kWh")
    print(f"最低购电费: {summary['total_cost_yuan']:.6f} 元")
    print(f"SOC范围: {summary['soc_min_kwh']:.6f} - {summary['soc_max_kwh']:.6f} kWh")
    print(f"最大单段充/放电量: {summary['max_charge_kwh']:.6f} / {summary['max_discharge_kwh']:.6f} kWh")
    print(f"平衡最大残差: {checks['balance_max_abs_kwh']:.12g} kWh")
    print(f"SOC递推最大残差: {checks['soc_update_max_abs_kwh']:.12g} kWh")
    print(f"同时充放电: {'是' if checks['simultaneous_charge_discharge'] else '否'}")
    print("表1购电量（kWh）:")
    for label, value in summary["table1"].items():
        print(f"  {label}: {value:.6f}")
    print("按模型起点连续24行聚合的充/放电量（kWh）:")
    for label, values in summary["positional_blocks"].items():
        print(f"  {label}: {values['charge_kwh']:.6f} / {values['discharge_kwh']:.6f}")
    print(f"边界SOC: {summary['soc_start_label']}={summary['soc_start_kwh']:.6f}, {summary['soc_end_label']}={summary['soc_end_kwh']:.6f}")
    print(f"计划购电标签与官方模板一致: {export_audit['purchase_template_labels_aligned']}")
    print(f"六段汇总标签与官方0:00-24:00表一致: {export_audit['official_storage_labels_aligned']}")
    print(f"候选Excel: {workbook_path}")
    print(f"逐时段审计CSV: {csv_path}\n")


def main() -> None:
    args = parse_arguments()
    paths = get_paths()

    # P0 occurs before the first model object or solver call is created.
    data, time_audit = audit_time_conventions(paths)
    print_p0_audit(time_audit)
    assert time_audit["completed_before_model"]
    assert not time_audit["unique_interpretation"]

    baseline_layout, baseline_values, template_hash, template_labels = template_baseline(paths.template_workbook)
    paths.output_dir.mkdir(parents=True, exist_ok=True)
    paths.formal_dir.mkdir(parents=True, exist_ok=True)
    paths.audit_dir.mkdir(parents=True, exist_ok=True)

    solutions: dict[str, Solution] = {}
    summaries: dict[str, dict[str, Any]] = {}
    numeric_solution = solve_milp(data, MAPPING_A)
    for mapping in (MAPPING_A, MAPPING_B):
        solution = replace(numeric_solution, mapping=mapping)
        checks = validate_numeric_solution(
            data,
            solution.purchase_kwh,
            solution.charge_kwh,
            solution.discharge_kwh,
            solution.curtailed_kwh,
            solution.charge_state,
            solution.soc_kwh[:-1],
            solution.soc_kwh[1:],
            MODEL_TOL,
        )
        summary = summarize(data, solution)
        assert abs(summary["total_cost_yuan"] - solution.objective_cost) <= MODEL_TOL
        assert checks["balance_max_abs_kwh"] <= MODEL_TOL
        solutions[mapping.code] = solution
        summaries[mapping.code] = summary

    comparison_path = paths.audit_dir / "time_mapping_comparison.csv"
    comparison = write_comparison(comparison_path, solutions, summaries)
    print("=== A/B比较 ===")
    for row in comparison.itertuples(index=False):
        print(f"{row.metric}: A={row.A:.6f}, B={row.B:.6f}, B-A={row.B_minus_A:.6f}")
    print(f"比较表: {comparison_path}")
    print("结论: 两种模式使用同一144行顺序并施加周期边界，故目标值相同；")
    print("表1按实际钟点取不同索引，SOC数值序列相同但绝对时间轴相差10分钟。\n")

    export_mappings = (
        (MAPPING_A, MAPPING_B)
        if args.time_mapping == "both"
        else (MAPPINGS[args.time_mapping],)
    )
    for selected_mapping in export_mappings:
        selected_solution = solutions[selected_mapping.code]
        selected_summary = summaries[selected_mapping.code]
        workbook_path, csv_path = mode_output_paths(paths, selected_mapping)
        solution_frame(data, selected_solution).to_csv(
            csv_path, index=False, encoding="utf-8-sig", float_format="%.12f"
        )
        write_candidate_workbook(paths, selected_solution, selected_summary, workbook_path)
        serialized, audited_data, serialized_checks = audit_serialized_schedule(
            csv_path, paths.input_workbook, selected_mapping, selected_solution.objective_cost
        )
        export_audit = audit_candidate_workbook(
            paths,
            selected_mapping,
            workbook_path,
            serialized,
            audited_data,
            selected_summary,
            baseline_layout,
            baseline_values,
            template_hash,
            template_labels,
        )
        assert serialized_checks["balance_max_abs_kwh"] <= MODEL_TOL
        assert export_audit["balance_max_abs_kwh"] <= MODEL_TOL
        print(f"=== Excel输出：模式{selected_mapping.code} ===")
        print_mode_report(
            selected_solution,
            selected_summary,
            serialized_checks,
            export_audit,
            workbook_path,
            csv_path,
        )

    # ``result1.xlsx`` is the official submission filename.  Keep the
    # natural-day/right-endpoint A solution as its canonical contents while
    # retaining the explicit A/B sensitivity outputs above.
    if any(mapping.code == "A" for mapping in export_mappings):
        official_result = paths.formal_dir / "result1.xlsx"
        shutil.copy2(paths.formal_dir / "result-A.xlsx", official_result)
        print(f"官方 result1 副本: {official_result}")

    assert sha256(paths.template_workbook) == template_hash
    print(
        f"自动检查均通过；本次输出模式{', '.join(item.code for item in export_mappings)}，"
        "原始附件和官方模板均未修改。"
    )


if __name__ == "__main__":
    main()
