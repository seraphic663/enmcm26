"""Export the final P2 Dynamic policy, including its interval evidence."""
import argparse, shutil
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from ..common import dataset as common
from ..common.intervals import write_storage_sheet, write_emergency_sheet
from ..common.paths import results_root

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--max-days',type=int)
    args=parser.parse_args()
    paths=common.get_paths()
    out=results_root()/'problem2/formal/dynamic'
    with np.load(out/'dynamic_intervals.npz', allow_pickle=False) as saved:
        arrays={key:saved[key].copy() for key in saved.files}
    records=pd.read_csv(out/'dynamic_daily.csv').to_dict('records')
    days=len(records)
    expected=args.max_days or 365
    if days != expected or len(arrays['dates']) != days:
        raise ValueError(f'Dynamic evidence length mismatch: expected {expected}, got {days}')
    if days<365:
        print(f'Dynamic smoke: {days} days; no formal workbook exported'); return
    dest=out/'result2.xlsx'; shutil.copy2(paths.template,dest)
    wb=load_workbook(dest)
    ws=wb['计划购电量']
    for t in range(144): ws.cell(1,t+2,common.interval_label(t))
    row_map={pd.Timestamp(ws.cell(r,1).value).strftime('%Y-%m-%d'):r for r in range(2,ws.max_row+1) if ws.cell(r,1).value is not None}
    for i,record in enumerate(records):
        if record['date'] not in row_map: continue
        r=row_map[record['date']]
        for t,value in enumerate(arrays['x'][i]): ws.cell(r,t+2,float(value))
        ws.cell(r,146,float(arrays['x'][i].sum())); ws.cell(r,147,record['planned_cost_yuan'])
    dates = [pd.Timestamp(str(date)) for date in arrays['dates']]
    out_idx = [i for i, date in enumerate(dates) if date >= common.OUTPUT_START]
    write_storage_sheet(wb, dates, out_idx, arrays['charge'], arrays['discharge'], arrays['soc'])
    write_emergency_sheet(wb, dates, out_idx, arrays['emergency'], 2, 1e-6)
    wb.save(dest); print(dest)
if __name__=='__main__': main()
