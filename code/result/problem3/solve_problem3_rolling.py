"""问题三滚动版（按 Problem3.md / 口径锁定）：在 baseline 副本上增加 6/12/18 调整。

保留不动：``solve_problem3_baseline.py``（无调整对照）。
本文件：策略 B 主结果 → 可写 result3.xlsx。

流程：
  0:00  Dynamic 负载 + 附件3 0点光伏 → 定 x_plan（与 baseline 同构）
  6/12/18  换光伏预报（负载保持 0:00）；禁止少购、冻储能、hold 门控
           目标含 c·x + 0.5c(|d+|+|d-|) + 5c·ê
  日终  冻结 x,u,v，口径 A 总费用 + 5c 紧急清算
"""

from __future__ import annotations
from mcm_code.common.paths import project_root, results_root

import argparse
import json
import shutil
from pathlib import Path

from ...common.intervals import date_rows, write_storage_sheet, write_emergency_sheet

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = project_root()

from ..problem2.solve_problem2_yesterday_dynamic import (
    paper_config,
    choose_candidate,
    choose_delta,
)
from ...common import dataset as common
from ...common import physical
from ...common.physical import rebuild_soc, settle_contract
from .problem3_adjust_milp import solve_stage
from .problem3_data import ISSUE_HOURS, issue_start_index, load_problem3_data

N = common.INTERVALS_PER_DAY
SOC_MIN = common.SOC_MIN
SOC_MAX = common.SOC_MAX
SOC_INITIAL = common.SOC_INITIAL
OUTPUT_START = common.OUTPUT_START
EPS_E = 1e-6

OUT_DIR = results_root() / "problem3" / "audit" / "rolling"
TEMPLATE = ROOT / "附件" / "附件5" / "result3.xlsx"
RESULT3 = results_root() / "problem3" / "formal" / "result3.xlsx"


def run_rolling(
    time_limit: float = 40.0,
    max_days: int | None = None,
) -> tuple[pd.DataFrame, dict, dict]:
    p3 = load_problem3_data(ROOT)
    data = common.from_arrays(p3.dates, p3.prices, p3.load_kwh, p3.pv_kwh)
    cold_load, cold_pv = p3.cold_load, p3.cold_pv
    pv_cube = p3.forecast_kwh
    config = paper_config()
    issue_hours = (6, 12, 18)
    allow_underbuy = False
    freeze_storage = True
    hold_gate = True
    load_adjustment = "hold"

    days = len(data.dates) if max_days is None else min(max_days, len(data.dates))
    records: list[dict] = []
    x_plans: list[np.ndarray] = []
    x_acts: list[np.ndarray] = []
    us: list[np.ndarray] = []
    vs: list[np.ndarray] = []
    es: list[np.ndarray] = []
    socs: list[np.ndarray] = []

    initial_soc = float(SOC_INITIAL)
    total_solver = 0.0
    bad_soc = bad_sim = 0

    for day in range(days):
        choice = choose_candidate(day, data, cold_load, cold_pv, config)
        delta_choice = choose_delta(
            day, choice.candidate, initial_soc, data, cold_load, cold_pv, config
        )
        delta = float(delta_choice.delta)
        load0 = choice.load + delta if config.reserve_in_milp else choice.load.copy()
        pv0 = pv_cube[day, 0, :]

        plan = physical.solve_physical_day(
            load0, pv0, data.prices, initial_soc, time_limit
        )
        day_solver = float(plan.solver_seconds)
        total_solver += day_solver

        x_plan = plan.x.copy() if config.reserve_in_milp else plan.x + delta
        x_act = x_plan.copy()
        charge = plan.charge.copy()
        discharge = plan.discharge.copy()
        soc = rebuild_soc(initial_soc, charge, discharge)
        bad_soc += int(
            plan.soc_residual_max > 1e-5
            or np.min(soc) < SOC_MIN - 1e-5
            or np.max(soc) > SOC_MAX + 1e-5
        )
        bad_sim += int(plan.simultaneous_count)

        for issue_hour in issue_hours:
            t0 = issue_start_index(issue_hour)
            issue_idx = ISSUE_HOURS.index(issue_hour)
            # 仅更新光伏预报，负载保持日初预测。
            load_tau = load0
            pv_tau = pv_cube[day, issue_idx, :]
            sol = solve_stage(
                t_start=t0,
                load_forecast=load_tau,
                pv_forecast=pv_tau,
                prices=data.prices,
                initial_soc=float(soc[t0]),
                x_plan=x_plan,
                locked_x=x_act,
                locked_charge=charge,
                locked_discharge=discharge,
                time_limit=time_limit,
                freeze_storage=freeze_storage,
            )
            day_solver += sol.solver_seconds
            total_solver += sol.solver_seconds
            # 预报域 hold vs 重优：仅在“同一预报下更优”时接受，避免无谓加购
            hold = settle_contract(
                load_tau, pv_tau, x_act, charge, discharge, data.prices, x_plan
            )
            trial = settle_contract(
                load_tau,
                pv_tau,
                sol.x,
                sol.charge,
                sol.discharge,
                data.prices,
                x_plan,
            )
            accept_trial = (
                float(trial["cost_total"]) + 1e-6
                < float(hold["cost_total"])
            )
            if accept_trial:
                x_act[t0:] = sol.x[t0:]
                charge[t0:] = sol.charge[t0:]
                discharge[t0:] = sol.discharge[t0:]
                soc = rebuild_soc(initial_soc, charge, discharge)
            bad_sim += int(((charge[t0:] > 1e-6) & (discharge[t0:] > 1e-6)).sum())

        settled = settle_contract(
            data.load_kwh[day],
            data.pv_kwh[day],
            x_act,
            charge,
            discharge,
            data.prices,
            x_plan,
        )
        emergency = np.asarray(settled["emergency"], dtype=float)

        records.append(
            {
                "date": str(data.dates[day].date()),
                "day_index": day,
                "selected_candidate": choice.candidate,
                "delta_kwh_per_interval": delta,
                "initial_soc_kwh": initial_soc,
                "terminal_soc_kwh": float(soc[-1]),
                "planned_purchase_kwh": float(x_plan.sum()),
                "actual_purchase_kwh": float(x_act.sum()),
                "planned_cost_ref_yuan": float(settled["cost_plan_ref"]),
                "contract_cost_yuan": float(settled["cost_contract"]),
                "actual_emergency_kwh": float(emergency.sum()),
                "actual_emergency_cost_yuan": float(settled["cost_emerg"]),
                "total_cost_yuan": float(settled["cost_total"]),
                "solver_seconds": day_solver,
            }
        )
        x_plans.append(x_plan)
        x_acts.append(x_act)
        us.append(charge)
        vs.append(discharge)
        es.append(emergency)
        socs.append(soc)

        initial_soc = float(soc[-1])
        if (day + 1) % 30 == 0 or day == 0:
            print(
                f"[P3-rolling] day {day + 1}/{days} {data.dates[day].date()} "
                f"cost={settled['cost_total']:.1f} "
                f"emerg={emergency.sum():.1f} "
                f"act-plan={x_act.sum() - x_plan.sum():+.1f} "
                f"soc_end={initial_soc:.1f}",
                flush=True,
            )

    daily = pd.DataFrame(records)
    output = daily[pd.to_datetime(daily["date"]) >= OUTPUT_START].copy()
    summary = {
        "name": "problem3_rolling_strategy_B",
        "description": (
            "baseline 0:00 plan + 6/12/18 adjust "
            "(PV-only, no underbuy, freeze storage, hold-gate); "
            "口径A; att3 multi-issue PV; noon-extrapolation fix"
        ),
        "config": config.name,
        "issue_hours": list(issue_hours),
        "allow_underbuy": bool(allow_underbuy),
        "freeze_storage": bool(freeze_storage),
        "hold_gate": bool(hold_gate),
        "load_adjustment": load_adjustment,
        "days_solved": int(days),
        "output_days": int(len(output)),
        "planned_purchase_kwh": float(output["planned_purchase_kwh"].sum()),
        "actual_purchase_kwh": float(output["actual_purchase_kwh"].sum()),
        "contract_cost_yuan": float(output["contract_cost_yuan"].sum()),
        "emergency_purchase_kwh": float(output["actual_emergency_kwh"].sum()),
        "emergency_cost_yuan": float(output["actual_emergency_cost_yuan"].sum()),
        "total_cost_yuan": float(output["total_cost_yuan"].sum()),
        "mean_daily_cost_yuan": float(output["total_cost_yuan"].mean()),
        "mean_delta_kwh": float(output["delta_kwh_per_interval"].mean()),
        "total_solver_seconds": float(total_solver),
        "bad_soc_flags": int(bad_soc),
        "simultaneous_intervals": int(bad_sim),
        "fixes": [
            "首小时用时效1外推（修复12:00正午抹零）",
            "调整禁止少购 allow_underbuy=False",
            "调整冻结储能 freeze_storage（只加购）",
            "调整不做上午负载偏差抬升（仅更新光伏）",
            "预报域 hold 门控：仅接受预报费用更低的调整",
        ],
        "note": "主结果可写入 result3.xlsx；baseline 文件未改动",
    }

    arrays = {
        "dates": np.array([r["date"] for r in records]),
        "x_plan": np.stack(x_plans),
        "x_act": np.stack(x_acts),
        "charge": np.stack(us),
        "discharge": np.stack(vs),
        "emergency": np.stack(es),
        "soc": np.stack(socs),
        "cost_total": np.array([r["total_cost_yuan"] for r in records]),
        "cost_plan_ref": np.array([r["planned_cost_ref_yuan"] for r in records]),
    }
    return daily, summary, arrays


def export_result3(arrays: dict, out_path: Path = RESULT3) -> Path:
    dates = [pd.Timestamp(str(d)) for d in arrays["dates"]]
    out_idx = [i for i, d in enumerate(dates) if d >= OUTPUT_START]
    if len(out_idx) != 334:
        raise ValueError(f"输出天数应为334，实际{len(out_idx)}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, out_path)
    wb = load_workbook(out_path)


    ws = wb["计划购电量"]
    date_row = date_rows(ws)
    for i in out_idx:
        r = date_row[dates[i].normalize()]
        xp = arrays["x_plan"][i]
        for t in range(N):
            ws.cell(r, t + 2, float(xp[t]))
        ws.cell(r, N + 2, float(xp.sum()))
        ws.cell(r, N + 3, float(arrays["cost_plan_ref"][i]))

    ws = wb["调整购电量"]
    date_row = date_rows(ws)
    for i in out_idx:
        r = date_row[dates[i].normalize()]
        xa = arrays["x_act"][i]
        for t in range(N):
            ws.cell(r, t + 2, float(xa[t]))
        ws.cell(r, N + 2, float(xa.sum()))
        ws.cell(r, N + 3, float(arrays["cost_total"][i]))

    write_storage_sheet(
        wb, dates, out_idx, arrays["charge"], arrays["discharge"], arrays["soc"]
    )
    write_emergency_sheet(wb, dates, out_idx, arrays["emergency"], 3, EPS_E)

    wb.save(out_path)
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--time-limit", type=float, default=40.0)
    parser.add_argument("--max-days", type=int, default=None)
    parser.add_argument("--no-export", action="store_true")
    args = parser.parse_args()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    daily, summary, arrays = run_rolling(args.time_limit, args.max_days)

    daily.to_csv(
        OUT_DIR / "problem3_rolling_daily.csv", index=False, encoding="utf-8-sig"
    )
    (OUT_DIR / "problem3_rolling_stats.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    np.savez_compressed(OUT_DIR / "problem3_rolling_daily.npz", **arrays)

    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)

    if not args.no_export and args.max_days is None:
        out = export_result3(arrays, RESULT3)
        print(f"result3 -> {out}", flush=True)
    elif not args.no_export:
        print("max-days 模式跳过 result3 导出", flush=True)

    print(f"rolling dir -> {OUT_DIR}", flush=True)


if __name__ == "__main__":
    main()
