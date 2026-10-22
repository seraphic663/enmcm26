"""Verify five reproduced workbooks and complete annual physical evidence.

Stored snapshots alone cannot pass. A different optimal schedule is reported as
a difference rather than silently accepted as a faithful numerical reproduction.
"""
import argparse
import json
from numbers import Real
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
from ..common.paths import SUPPORT_ROOT, protected_output_path
from ..common.dataset import get_paths, read_inputs

WORKBOOKS = ('result1.xlsx', 'result2.xlsx', 'result3.xlsx', 'result4-2.xlsx', 'result4-3.xlsx')
ARRAYS = (
    'problem2/formal/dynamic/dynamic_intervals.npz',
    'problem3/audit/baseline/problem3_baseline_daily.npz',
    'problem3/audit/rolling/problem3_rolling_daily.npz',
    *(f'problem4/audit/p4_2/p4_2_{tag}_arrays.npz' for tag in ('att1', 'yest', 'wd2')),
    *(f'problem4/audit/p4_3/p4_3_{tag}_arrays.npz' for tag in ('a_wd2', 'b0_yest', 'bstar_wd2', 'bfreeze_wd2')),
)
FIGURES = (
    '图01_微网购电与储能调度的整体建模框架.png',
    '图02_典型日负荷光伏电价净负荷.png',
    '图03_日前购电与储能联合调度结果.png',
    '图04_典型日期小区负载与光伏实际功率.png',
    '图05_不同购电策略间的总费用与紧急购电量对比.png',
    '图06_典型日光伏预测逐批更新.png',
)


def close(actual, expected, label, rtol=1e-8):
    a, b = np.asarray(actual, dtype=float), np.asarray(expected, dtype=float)
    if a.shape != b.shape or not np.isfinite(a).all() or not np.isfinite(b).all() or not np.allclose(a, b, atol=1e-5, rtol=rtol):
        raise ValueError(f'Mismatched {label}')


def verify_trace(a, load, pv, prices):
    """Independent physical and settlement equations; no solver/export helpers."""
    required = {'charge', 'discharge', 'soc', 'dates', 'emergency'}
    required |= {'x_plan', 'x_act'} if 'x_act' in a else {'x'}
    if not required.issubset(a):
        raise ValueError(f'Missing arrays: {sorted(required-set(a))}')
    x = np.asarray(a.get('x_act', a.get('x')), dtype=float)
    u, v, soc, emergency = (np.asarray(a[k], dtype=float) for k in ('charge', 'discharge', 'soc', 'emergency'))
    dates = np.asarray(a['dates']).astype(str)
    expected_dates = np.arange('2025-01-01', '2026-01-01', dtype='datetime64[D]').astype(str)
    if any(z.shape != (365,144) for z in (x,u,v,emergency)) or soc.shape != (365,145) or not np.array_equal(dates,expected_dates):
        raise ValueError('Incomplete annual dates or unexpected array shapes')
    for name in required-{'dates'}:
        if not np.isfinite(a[name]).all():
            raise ValueError(f'Nonfinite array: {name}')
    close(np.diff(soc, axis=1), .9*u-v/.9, 'SOC transitions', rtol=0)
    close(soc[1:,0], soc[:-1,-1], 'cross-day SOC', rtol=0)
    close(soc[0,0], 6000, 'initial SOC', rtol=0)
    simultaneous = int(np.count_nonzero((u > 1e-6)&(v > 1e-6)))
    if simultaneous or soc.min()<1200-1e-5 or soc.max()>10800+1e-5 or min(x.min(),u.min(),v.min(),emergency.min()) < -1e-5 or max(u.max(),v.max())>5000/6+1e-5:
        raise ValueError('Storage bounds, nonnegativity or mutual exclusion failed')
    gap = load+u-x-pv-v
    expected_emergency = np.maximum(gap,0)
    surplus = np.maximum(-gap,0)
    close(emergency,expected_emergency,'emergency purchases from actual load/PV',rtol=0)
    close(x+pv+v+emergency-load-u-surplus,np.zeros_like(x),'energy balance',rtol=0)
    plan = np.asarray(a.get('x_plan',x),dtype=float)
    if plan.shape != x.shape or not np.isfinite(plan).all() or plan.min() < -1e-5:
        raise ValueError('Invalid planned purchases')
    deviation = .5*prices*np.abs(x-plan) if 'x_act' in a else np.zeros_like(x)
    contract = np.sum(prices*x+deviation,axis=1)
    emergency_cost = np.sum(5*prices*expected_emergency,axis=1)
    metrics = dict(initial_soc_kwh=soc[:,0],terminal_soc_kwh=soc[:,-1],
        planned_purchase_kwh=plan.sum(axis=1),actual_purchase_kwh=x.sum(axis=1),
        planned_cost_yuan=np.sum(prices*plan,axis=1),contract_cost_yuan=contract,
        planned_cost_ref_yuan=np.sum(prices*plan,axis=1),plan_ref_cost_yuan=np.sum(prices*plan,axis=1),
        actual_emergency_kwh=expected_emergency.sum(axis=1),emergency_kwh=expected_emergency.sum(axis=1),
        actual_emergency_cost_yuan=emergency_cost,emergency_cost_yuan=emergency_cost,
        pv_curtailment_kwh=np.minimum(surplus,pv).sum(axis=1),
        grid_spill_kwh=np.maximum(surplus-pv,0).sum(axis=1),actual_surplus_kwh=surplus.sum(axis=1),
        total_cost_yuan=contract+emergency_cost)
    for name,target in [('cost_total',contract+emergency_cost),('cost_plan_ref',metrics['planned_cost_yuan']),('plan_ref_cost',metrics['planned_cost_yuan']),('planned_cost',metrics['planned_cost_yuan'])]:
        if name in a:
            close(a[name],target,name)
    return metrics, dict(status='PASS',days=365,intervals_per_day=144,
        soc_residual=float(np.max(np.abs(np.diff(soc,axis=1)-(.9*u-v/.9)))),
        cross_day_soc_residual=float(np.max(np.abs(soc[1:,0]-soc[:-1,-1]))),
        energy_balance_residual=float(np.max(np.abs(x+pv+v+emergency-load-u-surplus))),
        emergency_residual=float(np.max(np.abs(emergency-expected_emergency))),
        simultaneous_intervals=simultaneous,formal_total_cost_yuan=float((contract+emergency_cost)[31:].sum()))


def compare_workbook(actual, reference):
    if not actual.is_file() or not reference.is_file():
        raise ValueError(f'Missing workbook: {actual} or {reference}')
    got = load_workbook(actual, read_only=True, data_only=False)
    old = load_workbook(reference, read_only=True, data_only=False)
    result = dict(dimensions={}, numeric_cells=0, max_absolute_difference=0., different_cells=0, examples=[], excel_errors=[])
    try:
        if got.sheetnames != old.sheetnames:
            raise ValueError(f'Different worksheet names: {actual.name}')
        for ws in got:
            other = old[ws.title]
            result['dimensions'][ws.title] = [ws.max_row, ws.max_column]
            if (ws.max_row, ws.max_column) != (other.max_row, other.max_column):
                raise ValueError(f'Different worksheet dimensions: {actual.name}/{ws.title}')
            for row, original in zip(ws.iter_rows(), other.iter_rows()):
                for cell, reference_cell in zip(row, original):
                    value, target = cell.value, reference_cell.value
                    if cell.data_type == 'e':
                        result['excel_errors'].append(f'{ws.title}!{cell.coordinate}: {value}')
                    if isinstance(value, Real) and isinstance(target, Real):
                        result['numeric_cells'] += 1
                        result['max_absolute_difference'] = max(result['max_absolute_difference'], abs(float(value)-float(target)))
                        equal = bool(np.isfinite(value) and np.isfinite(target) and np.isclose(value, target, atol=1e-5, rtol=1e-8))
                    else:
                        equal = value == target
                    if not equal:
                        result['different_cells'] += 1
                        if len(result['examples']) < 12:
                            result['examples'].append(dict(cell=f'{ws.title}!{cell.coordinate}', actual=str(value), reference=str(target)))
        result['status'] = 'PASS' if not result['different_cells'] and not result['excel_errors'] else 'DIFFERENT'
        return result
    finally:
        got.close()
        old.close()


def verify_arrays(path, inputs, prices):
    if not path.is_file():
        raise ValueError(f'Missing required annual physical evidence: {path}')
    with np.load(path, allow_pickle=False) as saved:
        arrays = {k:saved[k] for k in saved.files}
    metrics, report = verify_trace(arrays, inputs.load_kwh, inputs.pv_kwh, prices)
    daily_path = path.with_name(path.name.replace('_arrays.npz','_daily.csv').replace('_daily.npz','_daily.csv').replace('dynamic_intervals.npz','dynamic_daily.csv'))
    daily = pd.read_csv(daily_path)
    if not np.array_equal(daily['date'].astype(str), arrays['dates'].astype(str)):
        raise ValueError(f'Daily CSV dates mismatch: {daily_path.name}')
    checked=[]
    for name, expected in metrics.items():
        if name in daily:
            close(daily[name], expected, f'{daily_path.name}/{name}')
            checked.append(name)
    must = {'initial_soc_kwh','terminal_soc_kwh','planned_purchase_kwh','total_cost_yuan'}
    must |= {'emergency_kwh','emergency_cost_yuan'} if 'problem4' in path.parts else {'actual_emergency_kwh','actual_emergency_cost_yuan'}
    if not must.issubset(checked):
        raise ValueError(f'Missing required daily settlement fields: {sorted(must-set(checked))}')
    report['daily_columns_checked']=checked
    return arrays, metrics, report


def verify_workbook_trace(path, a, metrics, threshold=1e-6):
    """Link annual interval evidence to exported dates, purchases, SOC and emergency blocks."""
    wb=load_workbook(path,read_only=True,data_only=True)
    dates=np.asarray(a['dates']).astype(str)
    formal=np.flatnonzero(dates>='2025-02-01')
    try:
        for title,key,cost in [('计划购电量','x_plan' if 'x_plan' in a else 'x','planned_cost_yuan'),('调整购电量','x_act','total_cost_yuan')]:
            if key not in a:
                continue
            rows=list(wb[title].values)[1:]
            if len(rows)!=334 or [pd.Timestamp(r[0]).strftime('%Y-%m-%d') for r in rows]!=dates[formal].tolist():
                raise ValueError(f'Workbook dates mismatch: {path.name}/{title}')
            close([r[1:145] for r in rows],a[key][formal],f'{path.name}/{title} intervals')
            close([r[145] for r in rows],a[key][formal].sum(axis=1),f'{path.name}/{title} totals')
            close([r[146] for r in rows],metrics[cost][formal],f'{path.name}/{title} costs')
        rows=list(wb['充放电量'].values)[1:]
        if len(rows)!=334*6:
            raise ValueError(f'Workbook storage block count: {path.name}')
        for j,i in enumerate(formal):
            chunk=rows[j*6:j*6+6]
            if pd.Timestamp(chunk[0][0]).strftime('%Y-%m-%d')!=dates[i]:
                raise ValueError(f'Workbook storage dates: {path.name}')
            close([r[2] for r in chunk],a['charge'][i].reshape(6,24).sum(axis=1),'storage charge blocks')
            close([r[3] for r in chunk],a['discharge'][i].reshape(6,24).sum(axis=1),'storage discharge blocks')
            close([chunk[0][5],chunk[1][5]],a['soc'][i,[0,-1]],'storage boundary SOC')
        # Parse displayed endpoints independently of the production interval merger.
        expected=[]
        def clock(index):
            return '0:00+1' if index==144 else f'{index//6}:{(index%6)*10:02d}'
        for i in formal:
            active=np.flatnonzero(a['emergency'][i]>threshold)
            for group in np.split(active,np.flatnonzero(np.diff(active)>1)+1):
                if group.size:
                    lo,hi=int(group[0]),int(group[-1])+1
                    expected.append((dates[i],f'{clock(lo)}-{clock(hi)}',a['emergency'][i,lo:hi].sum()))
        actual=[]; date=None
        for row in list(wb['紧急购电量'].values)[1:]:
            if row[0] is not None:
                date=pd.Timestamp(row[0]).strftime('%Y-%m-%d')
            actual.append((date,str(row[1]).replace('--','-'),row[2]))
        if [(r[0],r[1]) for r in actual]!=[(r[0],r[1]) for r in expected]:
            raise ValueError(f'Workbook emergency dates/intervals mismatch: {path.name}')
        close([r[2] for r in actual],[r[2] for r in expected],'workbook emergency amounts')
        return dict(status='PASS',formal_days=334,storage_blocks=len(rows),emergency_blocks=len(actual))
    finally:
        wb.close()


def verify_figures(directory):
    actual={p.name for p in directory.glob('*.png')}
    if actual!=set(FIGURES):
        raise ValueError(f'Figure names mismatch; missing={sorted(set(FIGURES)-actual)}, extra={sorted(actual-set(FIGURES))}')
    report={}
    for name in FIGURES:
        path=directory/name
        with Image.open(path) as im:
            if im.format!='PNG' or min(im.size)<1:
                raise ValueError(f'Invalid PNG: {name}')
            im.verify()
        with Image.open(path) as im:
            im.load()
            report[name]=dict(status='PASS',width=im.width,height=im.height)
    return report


def verify_perfect_information(evidence, inputs):
    frame=pd.read_csv(evidence/'problem2/audit/problem2_A_schedule.csv')
    if len(frame)!=365*144:
        raise ValueError('Incomplete P2 A schedule')
    dates=frame['date'].to_numpy().reshape(365,144)
    if not np.array_equal(dates,np.repeat(inputs.dates.strftime('%Y-%m-%d').to_numpy()[:,None],144,axis=1)):
        raise ValueError('P2 A schedule date ordering')
    close(frame['interval_index'],np.tile(np.arange(144),365),'P2 A interval ordering')
    def values(key): return frame[key].to_numpy().reshape(365,144)
    close(values('load_kwh'),inputs.load_kwh,'P2 A source load')
    close(values('pv_kwh'),inputs.pv_kwh,'P2 A source PV')
    close(values('price_yuan_per_kwh'),np.broadcast_to(inputs.prices,(365,144)),'P2 A source prices')
    a=dict(dates=dates[:,0],x=values('planned_purchase_kwh'),charge=values('charge_kwh'),discharge=values('discharge_kwh'),emergency=values('emergency_purchase_kwh'),soc=np.column_stack([values('soc_start_kwh')[:,0],values('soc_end_kwh')]))
    close(values('soc_start_kwh'),a['soc'][:,:-1],'P2 A SOC start/end chain')
    metrics,report=verify_trace(a,inputs.load_kwh,inputs.pv_kwh,inputs.prices)
    audit=json.loads((evidence/'problem2/audit/problem2_A_audit.json').read_text(encoding='utf-8'))
    if audit['solver']['status']!='Optimal' or audit['solver']['mip_gap']!=0:
        raise ValueError('P2 A optimality not proven')
    close(audit['solver']['objective_full_year_yuan'],metrics['total_cost_yuan'].sum(),'P2 A annual objective')
    close(audit['output_period']['total_purchase_cost_yuan'],metrics['total_cost_yuan'][31:].sum(),'P2 A formal cost')
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, required=True, help='Completed annual reproduction output')
    parser.add_argument('--reference', type=Path, default=SUPPORT_ROOT)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    root, reference = args.root.resolve(), args.reference.resolve()
    if args.report:
        args.report=protected_output_path(args.report)
    report = dict(workbooks={}, arrays={}, workbook_links={}, errors=[], tolerances=dict(workbook_atol=1e-5, workbook_rtol=1e-8, physical_kwh=1e-5))
    inputs=read_inputs(get_paths())
    price_frame=pd.read_excel(SUPPORT_ROOT/'data/附件/附件4.xlsx')
    if not pd.DatetimeIndex(pd.to_datetime(price_frame.iloc[:,0])).equals(inputs.dates):
        raise ValueError('Actual price dates mismatch')
    actual_prices=price_frame.iloc[:,1:].to_numpy(dtype=float)
    if actual_prices.shape!=(365,144) or not np.isfinite(actual_prices).all() or actual_prices.min()<=0:
        raise ValueError('Invalid actual price input')
    evidence=root/'result/来源证据'
    for name in WORKBOOKS:
        try:
            report['workbooks'][name] = compare_workbook(root/'result'/name, reference/'result'/name)
        except (OSError, ValueError) as exc:
            report['errors'].append(str(exc))
    link_names={ARRAYS[0]:'result2.xlsx',ARRAYS[2]:'result3.xlsx',ARRAYS[5]:'result4-2.xlsx',ARRAYS[9]:'result4-3.xlsx'}
    for name in ARRAYS:
        try:
            a,metrics,report['arrays'][name] = verify_arrays(evidence/name,inputs,actual_prices if name.startswith('problem4/') else inputs.prices)
            if name in link_names:
                workbook=link_names[name]
                report['workbook_links'][workbook]=verify_workbook_trace(root/'result'/workbook,a,metrics,1e-7 if name.startswith('problem4/') else 1e-6)
        except (OSError, ValueError, KeyError) as exc:
            report['errors'].append(f'{name}: {exc}')
    try:
        report['perfect_information']=verify_perfect_information(evidence,inputs)
    except (OSError,ValueError,KeyError) as exc:
        report['errors'].append(f'P2 A: {exc}')
    try:
        report['figures']=verify_figures(root/'figure')
    except (OSError,ValueError) as exc:
        report['errors'].append(str(exc))
    report['status'] = 'PASS' if not report['errors'] and all(x['status'] == 'PASS' for x in report['workbooks'].values()) else 'FAIL'
    report['scope'] = 'Five frozen workbook comparisons; ten full-year traces independently checked against raw load/PV/prices and daily costs; four annual workbook links; P2 A physics and optimality; six named decodable PNGs. Tables are checked separately.'
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(text+'\n', encoding='utf-8')
    print(text)
    if report['status'] != 'PASS':
        raise SystemExit(1)


if __name__ == '__main__':
    main()
